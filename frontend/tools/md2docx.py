import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x8B, 0x2E, 0x1F)   # đỏ son
INK = RGBColor(0x22, 0x22, 0x22)
CODE_FONT = "Consolas"
BODY_FONT = "Calibri"


def shade(el, hex_fill):
    pr = el.get_or_add_tcPr() if el.tag.endswith("tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def left_border(p, color="8B2E1F"):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    l = OxmlElement("w:left")
    for k, v in {"w:val": "single", "w:sz": "18", "w:space": "8", "w:color": color}.items():
        l.set(qn(k), v)
    bdr.append(l)
    pPr.append(bdr)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<!\*)\*[^*\s][^*]*\*(?!\*))")


def add_inline(p, text, size=None, bold=False, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            add_inline(p, part[2:-2], size=size, bold=True, color=color)
            continue
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = CODE_FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
            r.font.color.rgb = ACCENT
            if size: r.font.size = Pt(size - 0.5)
            if bold: r.bold = True
            continue
        elif part.startswith("[") and "](" in part:
            label, url = part[1:part.index("](")], part[part.index("](") + 2:-1]
            r = p.add_run(label); r.underline = True
            if url.startswith("http"):
                r2 = p.add_run(f" ({url})"); r2.font.size = Pt((size or 10.5) - 1.5)
                r2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        if size: r.font.size = Pt(size)
        if bold: r.bold = True
        if color: r.font.color.rgb = color


def code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.3)
    shade(p._p, "F4F1EA")
    for i, ln in enumerate(lines):
        r = p.add_run(ln if ln else " ")
        r.font.name = CODE_FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
        r._element.rPr.rFonts.set(qn("w:cs"), CODE_FONT)
        r.font.size = Pt(8.5)
        if i < len(lines) - 1:
            r.add_break()


NODE = re.compile(r'^\s*(\w+)(?:\[\(?"(.+?)"\)?\]|\{"(.+?)"\})?(?::::\w+)?\s*$')


def mermaid_to_list(doc, lines):
    labels, edges = {}, []
    def node(tok):
        m = NODE.match(tok)
        if not m:
            return tok.strip()
        nid, a, b = m.groups()
        lab = a or b
        if lab:
            labels[nid] = lab.replace("<br/>", " – ")
        return nid
    for ln in lines:
        if ln.strip().startswith(("flowchart", "classDef", "graph")):
            continue
        m = re.match(r"\s*(.+?)\s*(?:--\s*(.+?)\s*-->|-\.->|-->)\s*(.+)$", ln)
        if not m:
            continue
        a, lab, b = m.groups()
        dashed = "-.->" in ln
        edges.append((node(a), node(b), lab.strip('"') if lab else lab, dashed))
    p = doc.add_paragraph()
    add_inline(p, "**Sơ đồ luồng (dạng danh sách):**", size=10.5)
    for i, (a, b, lab, dashed) in enumerate(edges, 1):
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(0.8)
        q.paragraph_format.first_line_indent = Cm(-0.6)
        q.paragraph_format.space_after = Pt(2)
        txt = f"{i}. {labels.get(a, a)}  →  {labels.get(b, b)}"
        if lab:
            txt += f"  (khi: {lab})"
        if dashed:
            txt += "  (nhánh dự phòng)"
        add_inline(q, txt, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[2:]]
    n = len(header)
    body = [r + [""] * (n - len(r)) for r in body]
    t = doc.add_table(rows=1 + len(body), cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    total = 17.0
    weights = [max(6, max(len(r[i]) for r in [header] + body)) for i in range(n)]
    weights = [min(w, 70) for w in weights]
    s = sum(weights)
    widths = [Cm(total * w / s) for w in weights]
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for ri, row in enumerate([header] + body):
        for ci, val in enumerate(row[:n]):
            cell = t.cell(ri, ci)
            cell.width = widths[ci]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            if ri == 0:
                p.paragraph_format.keep_with_next = True
                shade(cell._tc, "8B2E1F")
                add_inline(p, val, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            else:
                if ri % 2 == 0:
                    shade(cell._tc, "FAF6EF")
                add_inline(p, val, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(src, dst):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2))
    sec.top_margin = sec.bottom_margin = Cm(2)

    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        h = doc.styles[name]
        h.font.name = BODY_FONT
        h.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = ACCENT if name in ("Title", "Heading 1") else INK
        h.paragraph_format.space_before = Pt(14 if name != "Heading 3" else 10)
        h.paragraph_format.space_after = Pt(6)

    # footer page number
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        r._r.append(e)
    r.font.size = Pt(9)

    lines = open(src, encoding="utf-8").read().splitlines()
    i, first_h1 = 0, True
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("```"):
            lang = s[3:].strip()
            j = i + 1
            block = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                block.append(lines[j]); j += 1
            if lang == "mermaid":
                mermaid_to_list(doc, block)
            else:
                code_block(doc, block)
            i = j + 1
            continue
        if s.startswith("|"):
            j = i
            rows = []
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(lines[j]); j += 1
            table(doc, rows)
            i = j
            continue
        if not s:
            i += 1; continue
        if s.startswith("# "):
            if first_h1:
                p = doc.add_paragraph(style="Title"); add_inline(p, s[2:]); first_h1 = False
            else:
                doc.add_heading(level=1).add_run(s[2:])
        elif s.startswith("### "):
            p = doc.add_heading(level=3); add_inline(p, s[4:])
        elif s.startswith("## "):
            p = doc.add_heading(level=1); add_inline(p, s[3:])
        elif s.startswith(">"):
            text = s.lstrip("> ").strip()
            p = doc.add_paragraph()
            left_border(p)
            p.paragraph_format.left_indent = Cm(0.4)
            add_inline(p, text, size=10, color=RGBColor(0x55, 0x55, 0x55))
        elif re.match(r"^- \[[ xX]\] ", s):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, "☐  " + s[6:])
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, s[2:])
        elif re.match(r"^\d+\. ", s):
            num, rest = s.split(". ", 1)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, f"{num}.  " + rest)
        else:
            p = doc.add_paragraph(); add_inline(p, s)
        i += 1
    doc.save(dst)


if __name__ == "__main__":
    for a in sys.argv[1:]:
        convert(a, a[:-3] + ".docx")
        print("ok", a[:-3] + ".docx")
